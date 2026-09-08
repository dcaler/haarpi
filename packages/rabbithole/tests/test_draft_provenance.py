"""What a redlined draft says about itself: its stats, its bibliography structure, its log.

Three ways a reviewed document misdescribed itself. The redline path carried the front-matter
through untouched, so an elephantRoom draft claimed "Sources: 162" directly above a block that
had just recounted them and said 184, under an eight-week-old date. The bibliography's tier
headings came back as bold paragraphs — identical to the 300 bold citation lines they were
meant to divide, and absent from the navigation pane. And nothing wrote a durable run log, so
the reasoning behind a ten-hour run aged out of a 5.8KB trundlr tail.

Runnable two ways:
    pytest tests/test_draft_provenance.py
    python tests/test_draft_provenance.py
"""

from __future__ import annotations

from pathlib import Path

import pytest

from rabbithole import redline

docx = pytest.importorskip("docx")


def _doc(tmp_path: Path, paras: list[tuple[str, str]]) -> Path:
    from docx import Document
    d = Document()
    for style, text in paras:
        p = d.add_paragraph()
        p.add_run(text)
        if style:
            p.style = style
    fp = tmp_path / "draft.docx"
    d.save(str(fp))
    return fp


def _read(fp: Path):
    from docx import Document
    return [(p.style.name if p.style is not None else "", p.text)
            for p in Document(str(fp)).paragraphs]


# ── front matter describes THIS draft ────────────────────────────────────────

def test_the_front_matter_stats_are_rewritten(tmp_path):
    fp = _doc(tmp_path, [("Title", "Literature Review"),
                         ("", "Project: p\nDate: 2026-08-15\nSources: 162")])
    assert redline.replace_front_matter(
        fp, [("Project:", "p"), ("Date:", "2026-09-08"), ("Sources:", "184")]) == {"front_matter": 1}
    text = _read(fp)[1][1]
    assert "184" in text and "2026-09-08" in text
    assert "162" not in text and "2026-08-15" not in text, "a stale stat is not left standing"


def test_a_draft_with_no_front_matter_is_left_alone(tmp_path):
    fp = _doc(tmp_path, [("Title", "Literature Review"), ("", "Straight into the prose.")])
    assert redline.replace_front_matter(fp, [("Project:", "p")]) == {"front_matter": 0}
    assert _read(fp)[1][1] == "Straight into the prose."


def test_the_labels_stay_italic_and_the_values_do_not(tmp_path):
    """Cosmetic, but it is how the block reads as a block rather than a run-on line."""
    from docx import Document
    fp = _doc(tmp_path, [("Title", "T"), ("", "Project: old")])
    redline.replace_front_matter(fp, [("Project:", "new"), ("Date:", "2026-09-08")])
    runs = Document(str(fp)).paragraphs[1].runs
    labelled = {r.text.strip(): bool(r.font.italic) for r in runs if r.text.strip()}
    assert labelled["Project:"] is True and labelled["Date:"] is True
    assert labelled["new"] is False and labelled["2026-09-08"] is False


# ── the bibliography keeps its structure through a redline ───────────────────

_BIB = ("## Annotated Bibliography\n\n"
        "### Cited in the review\n\n**Smith, A. (2020). A paper.**\n\n- a claim\n\n"
        "### Additional curated sources\n\n**Jones, B. (2021). Another.**\n\n- another claim\n")


def test_bibliography_tier_headings_are_real_headings(tmp_path):
    """They were bold paragraphs, and every citation line below them is also bold — so the
    division was invisible, and neither heading reached the navigation pane."""
    fp = _doc(tmp_path, [("Title", "T"), ("", "Body prose."),
                         ("Heading 2", "Annotated Bibliography"), ("", "old entry")])
    redline.replace_bibliography(fp, _BIB)
    styles = {t: s for s, t in _read(fp)}
    assert styles["Cited in the review"] == "Heading 3"
    assert styles["Additional curated sources"] == "Heading 3"


def test_the_entries_themselves_are_not_headings(tmp_path):
    """The fix must divide the list, not promote every line in it."""
    fp = _doc(tmp_path, [("Title", "T"), ("Heading 2", "Annotated Bibliography")])
    redline.replace_bibliography(fp, _BIB)
    styles = {t: s for s, t in _read(fp)}
    assert styles["Smith, A. (2020). A paper."] != "Heading 3"


def test_body_above_the_bibliography_is_untouched(tmp_path):
    fp = _doc(tmp_path, [("Title", "T"), ("", "Body prose that must survive."),
                         ("Heading 2", "Annotated Bibliography"), ("", "old entry")])
    redline.replace_bibliography(fp, _BIB)
    assert _read(fp)[1][1] == "Body prose that must survive."


def test_a_missing_heading_style_falls_back_to_bold_not_to_nothing(tmp_path):
    """A style id absent from the document renders as body text, which is worse than the bold
    it replaced — so the fallback is the old behaviour, never none."""
    from docx import Document
    fp = _doc(tmp_path, [("Title", "T")])
    d = Document(str(fp))
    p = d.add_paragraph()
    p.add_run("Cited in the review")
    assert redline._apply_style(p, "No Such Style ZZZ") is False
    assert p.runs[0].bold is True


# ── a run leaves a record ────────────────────────────────────────────────────

def test_a_run_writes_a_durable_log(tmp_path, capsys):
    """`.haarpi/runlog/` was created by `init` and never written to; the only record of a
    ten-hour verb was a few kilobytes of trundlr tail."""
    from haarpi import runlog
    import sys
    saved_out, saved_err = sys.stdout, sys.stderr
    try:
        fp = runlog.to_file(tmp_path, "rabbithole_revise")
        print("a line the log must keep")
    finally:
        sys.stdout, sys.stderr = saved_out, saved_err
    assert fp is not None and fp.exists()
    assert fp.parent == tmp_path / ".haarpi" / "runlog"
    assert "revise" in fp.name
    assert "a line the log must keep" in fp.read_text()


def test_an_unwritable_log_does_not_stop_the_run(tmp_path):
    """Best-effort: a log that cannot be opened is not a reason to refuse to do the work."""
    from haarpi import runlog
    blocker = tmp_path / ".haarpi"
    blocker.write_text("not a directory")
    assert runlog.to_file(tmp_path, "gather") is None


if __name__ == "__main__":
    import tempfile, traceback
    fns = [v for k, v in sorted(globals().items())
           if k.startswith("test_") and callable(v)]
    failures = 0
    for fn in fns:
        try:
            n = fn.__code__.co_argcount
            with tempfile.TemporaryDirectory() as td:
                fn(*([Path(td)] + ([None] if n > 1 else []))[:n])
            print(f"  PASS  {fn.__name__}")
        except Exception:  # noqa: BLE001
            failures += 1
            print(f"  FAIL  {fn.__name__}")
            traceback.print_exc()
    print(f"\n{len(fns) - failures}/{len(fns)} passed")
    raise SystemExit(1 if failures else 0)
