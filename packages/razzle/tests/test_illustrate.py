"""razzle.illustrate — pictorial briefs become prompts, `imagine` draws candidates, the AUTHOR picks,
and only a pick reaches a slide. Schematic briefs are never drawn."""

from __future__ import annotations

import json
import os
import stat

import pytest

from razzle import assets, compose, illustrate, render

_MASTER = assets.master_pptx("default")
_DESC = assets.descriptor("default")


def _spec():
    return [{"role": "title", "title": "A talk"},
            {"role": "content", "title": "Crowds sort themselves", "body": ["a"],
             "illustration": "a crowded street market seen from above",
             "illustration_kind": "pictorial"},
            {"role": "content", "title": "The rule", "body": ["b"],
             "illustration": "a diagram of the update rule", "illustration_kind": "pictorial"},
            {"role": "content", "title": "Takeaway", "body": ["c"]}]


# ── kinds ────────────────────────────────────────────────────────────────────

def test_a_brief_naming_a_diagram_is_schematic_whatever_it_was_tagged():
    assert illustrate.kind_of({"illustration": "a bar chart of the sweep",
                               "illustration_kind": "pictorial"}) == "schematic"
    assert illustrate.kind_of({"illustration": "a lighthouse in fog"}) == "pictorial"
    assert illustrate.kind_of({"illustration": "a lighthouse in fog",
                               "illustration_kind": "schematic"}) == "schematic"


def test_normalise_keeps_the_kind_only_beside_a_brief():
    slides = compose.normalise([
        {"role": "title", "title": "T"},
        {"role": "content", "title": "A", "illustration": "a lighthouse in fog",
         "illustration_kind": "pictorial"},
        {"role": "content", "title": "B", "illustration_kind": "pictorial"},
        {"role": "content", "title": "C", "illustration": "a flow chart of the method"}], set())
    assert slides[1]["illustration_kind"] == "pictorial"
    assert "illustration_kind" not in slides[2]
    assert slides[3]["illustration_kind"] == "schematic"


def test_only_pictorial_briefs_become_prompts():
    bs = illustrate.briefs(_spec())
    assert [b["slide"] for b in bs] == [2]
    assert bs[0]["prompt"].startswith("a crowded street market seen from above.")
    assert "no text" in bs[0]["prompt"]


def test_a_key_survives_a_reauthor_that_keeps_the_brief_and_changes_with_it():
    k = illustrate.key_for(1, "a lighthouse in fog")
    assert k == illustrate.key_for(1, "a lighthouse in fog")
    assert k != illustrate.key_for(1, "a lighthouse at dusk")


# ── the render's half ────────────────────────────────────────────────────────

def _pick(deck_dir, key, data=b"png-bytes"):
    d = illustrate.home(deck_dir) / key
    d.mkdir(parents=True, exist_ok=True)
    (d / "cand_0.png").write_bytes(data)
    (d / illustrate.CHOSEN).write_bytes(data)


def test_prepare_writes_the_prompt_list_and_a_note_while_anything_is_unpicked(tmp_path):
    ill = illustrate.prepare(tmp_path, _spec())
    assert json.loads((tmp_path / "illustrations" / "prompts.json").read_text())[0]["slide"] == 2
    assert ill["placements"] == {}
    assert "Not yet picked: slide 2" in ill["note"]
    assert "UNRESOLVED" in ill["note"]


def test_a_deck_with_no_pictorial_brief_gets_no_note_and_no_folder(tmp_path):
    spec = [{"role": "content", "title": "T", "illustration": "a schematic of the model"}]
    ill = illustrate.prepare(tmp_path, spec)
    assert ill["note"] is None and ill["briefs"] == []
    assert not (tmp_path / "illustrations").exists()


def test_a_pick_is_placed_and_stays_placed_until_it_changes(tmp_path):
    spec = _spec()
    key = illustrate.briefs(spec)[0]["key"]
    illustrate.prepare(tmp_path, spec)
    _pick(tmp_path, key)
    assert illustrate.needs_placing(tmp_path) == [key]           # chosen, deck not re-rendered
    ill = illustrate.prepare(tmp_path, spec)
    assert set(ill["placements"]) == {1}
    assert "Placed from your picks: slide 2" in ill["note"]
    illustrate.record_placed(tmp_path, ill["placements"], ill["briefs"])
    assert illustrate.needs_placing(tmp_path) == []
    _pick(tmp_path, key, b"a different candidate")               # the author changed their mind
    assert illustrate.needs_placing(tmp_path) == [key]


# ── the verb ─────────────────────────────────────────────────────────────────

def _fake_imagine(bin_dir, *, busy_first: int = 0):
    """An `imagine` that writes -n files named like the real one, refusing `busy_first` times."""
    counter = bin_dir / "calls"
    counter.write_text("0")
    script = bin_dir / "imagine"
    script.write_text(f"""#!/bin/bash
n=1; out=""
while [ $# -gt 0 ]; do case "$1" in -n) n=$2; shift 2;; -o) out=$2; shift 2;; -s) shift 2;; *) shift;; esac; done
c=$(cat {counter}); echo $((c+1)) > {counter}
if [ "$c" -lt {busy_first} ]; then echo "imagine: the freest card, GPU 0, has 900 MiB free; need 13000."; exit 1; fi
if [ "$n" -eq 1 ]; then echo x > "$out.png"; else for i in $(seq 0 $((n-1))); do echo x > "${{out}}_$i.png"; done; fi
""")
    script.chmod(script.stat().st_mode | stat.S_IEXEC)
    return counter


@pytest.fixture
def on_path(tmp_path, monkeypatch):
    bin_dir = tmp_path / "bin"
    bin_dir.mkdir()
    monkeypatch.setenv("PATH", f"{bin_dir}{os.pathsep}{os.environ['PATH']}")
    return bin_dir


def test_illustrate_draws_each_brief_once(tmp_path, on_path):
    calls = _fake_imagine(on_path)
    deck = tmp_path / "deck"
    illustrate.prepare(deck, _spec())
    assert illustrate.run(deck, count=3) == 0
    key = illustrate.read_prompts(deck)[0]["key"]
    assert len(illustrate.candidates(deck, key)) == 3
    assert illustrate.run(deck, count=3) == 0                    # nothing left to draw
    assert calls.read_text().strip() == "1"


def test_illustrate_waits_out_a_card_that_is_still_busy(tmp_path, on_path):
    calls = _fake_imagine(on_path, busy_first=2)
    deck = tmp_path / "deck"
    illustrate.prepare(deck, _spec())
    assert illustrate.run(deck, count=2, wait_s=0) == 0
    assert calls.read_text().strip() == "3"


def test_illustrate_fails_loudly_without_imagine(tmp_path, monkeypatch):
    monkeypatch.setattr(illustrate.shutil, "which", lambda *_: None)
    illustrate.prepare(tmp_path, _spec())
    assert illustrate.run(tmp_path) == 1


def test_illustrate_with_nothing_to_draw_is_a_clean_no_op(tmp_path):
    assert illustrate.run(tmp_path) == 0


# ── on the slide ─────────────────────────────────────────────────────────────

@pytest.mark.skipif(not (_MASTER and _DESC), reason="neutral house master absent")
def test_a_picked_illustration_renders_beside_the_bullets_and_says_where_it_came_from(tmp_path):
    from PIL import Image
    img = tmp_path / "pick.png"
    Image.new("RGB", (64, 64), (200, 120, 40)).save(img)
    spec = _spec()
    out = render.render_deck(spec, _DESC["master_path"], _DESC, tmp_path / "d.pptx",
                             illustrations={1: img})
    from pptx import Presentation
    prs = Presentation(str(out))
    s = prs.slides[1]
    assert any(sh.shape_type == 13 for sh in s.shapes)           # 13 == PICTURE
    assert "a" in " ".join(sh.text_frame.text for sh in s.shapes if sh.has_text_frame)
    assert s.notes_slide.notes_text_frame.text.startswith("ILLUSTRATION (generated): ")
    # the schematic brief on slide 3 is untouched: still a TODO, nothing drawn
    assert prs.slides[2].notes_slide.notes_text_frame.text == (
        "ILLUSTRATION: a diagram of the update rule")
    assert not any(sh.shape_type == 13 for sh in prs.slides[2].shapes)


@pytest.mark.skipif(not (_MASTER and _DESC), reason="neutral house master absent")
def test_razzles_note_is_an_unresolved_comment_signed_by_the_tool(tmp_path):
    from haarpi import redline
    out = render.render_deck(_spec(), _DESC["master_path"], _DESC, tmp_path / "d.pptx",
                             note="Illustrations pending.", note_author=("razzle", "ra"))
    threads = redline.pptx_comment_threads(out)
    assert [(t["author"], t["text"], t["resolved"], t["slide"]) for t in threads] == [
        ("razzle", "Illustrations pending.", False, "slide1")]
    assert redline.is_tool_author("razzle")
    assert not redline.gate_check(out)["clean"]                  # it holds the mint
    from pptx import Presentation
    Presentation(str(out))                                       # and the package still opens
